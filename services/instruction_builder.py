from __future__ import annotations

from datetime import datetime
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from services.instruction_naming import instruction_title


TEMPLATE_PATH = Path(
    "templates/ИОТ ИСХОДНИК.docx"
)

DEFAULT_OUTPUT_DIR = Path(
    "output"
)

FONT_NAME = "Times New Roman"
BODY_FONT_SIZE = 10
TITLE_FONT_SIZE = 12

SECTION_HEADING_RE = re.compile(
    r"^\s*\d+\.\s+\S"
)

POINT_RE = re.compile(
    r"^\s*(\d+\.\d+)\.\s*(.*)$"
)


def _set_run_font(
    run,
    font_name=FONT_NAME,
    font_size=BODY_FONT_SIZE,
):
    run.font.name = font_name
    run.font.size = Pt(
        font_size
    )

    r_fonts = run._element.rPr.rFonts

    r_fonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii",
        font_name,
    )

    r_fonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi",
        font_name,
    )

    r_fonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cs",
        font_name,
    )


def _set_paragraph_font(
    paragraph,
    font_name=FONT_NAME,
    font_size=BODY_FONT_SIZE,
):
    for run in paragraph.runs:
        _set_run_font(
            run,
            font_name=font_name,
            font_size=font_size,
        )


def _replace_paragraph_text(
    paragraph,
    text: str,
    font_size=BODY_FONT_SIZE,
):
    paragraph.clear()

    run = paragraph.add_run(
        text
    )

    _set_run_font(
        run,
        font_size=font_size,
    )


def _clean_generated_text(
    text: str,
) -> str:
    text = str(
        text or ""
    )

    replacements = {
        "**": "",
        "__": "",
        "###": "",
        "##": "",
        "#": "",
    }

    for old, new in replacements.items():
        text = text.replace(
            old,
            new,
        )

    text = text.replace(
        "\r\n",
        "\n",
    )

    text = text.replace(
        "\r",
        "\n",
    )

    lines = [
        line.strip()
        for line in text.split(
            "\n"
        )
        if line.strip()
    ]

    return "\n".join(
        lines
    )


def split_instruction_text(
    instruction_text: str,
) -> list[str]:
    """
    Возвращает строки готовой ИОТ.

    Генератор уже обязан вернуть нумерацию вида:
      1. ОБЩИЕ...
      1.1. ...
      1.2. ...
      2. ТРЕБОВАНИЯ...

    Здесь не пытаемся менять смысл текста.
    """
    cleaned = _clean_generated_text(
        instruction_text
    )

    return [
        line.strip()
        for line in cleaned.split(
            "\n"
        )
        if line.strip()
    ]


def _find_body_start_index(
    document: Document,
) -> int | None:
    """
    Ищет начало старого текста ИОТ в шаблоне.
    """
    for index, paragraph in enumerate(
        document.paragraphs
    ):
        text = " ".join(
            paragraph.text.split()
        ).strip()

        if text.startswith(
            "1. ОБЩИЕ"
        ):
            return index

    return None


def _remove_old_body(
    document: Document,
) -> None:
    start_index = _find_body_start_index(
        document
    )

    if start_index is None:
        return

    paragraphs = document.paragraphs

    for paragraph in paragraphs[
        start_index:
    ]:
        element = paragraph._element
        element.getparent().remove(
            element
        )


def _find_title_paragraphs(
    document: Document,
):
    instruction_heading = None
    subtitle = None

    for paragraph in document.paragraphs:
        text = " ".join(
            paragraph.text.split()
        ).strip()

        if (
            instruction_heading is None
            and text.upper()
            == "ИНСТРУКЦИЯ"
        ):
            instruction_heading = paragraph
            continue

        if (
            subtitle is None
            and text.lower().startswith(
                "по охране труда "
            )
        ):
            subtitle = paragraph

    return (
        instruction_heading,
        subtitle,
    )


def _set_document_title(
    document: Document,
    target_name: str,
    target_kind: str,
):
    """
    В отдельной ИОТ объект и адрес остаются плейсхолдерами.
    Они будут заменены только ручным VBA ReplaceTags
    после проверки итогового DOCM.
    """
    instruction_heading, subtitle = (
        _find_title_paragraphs(
            document
        )
    )

    if instruction_heading is None:
        raise RuntimeError(
            "В шаблоне ИОТ не найден заголовок «ИНСТРУКЦИЯ»."
        )

    if subtitle is None:
        raise RuntimeError(
            "В шаблоне ИОТ не найден заголовок "
            "«по охране труда ...»."
        )

    _replace_paragraph_text(
        instruction_heading,
        "ИНСТРУКЦИЯ",
        font_size=TITLE_FONT_SIZE,
    )

    instruction_heading.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    instruction_heading.paragraph_format.space_before = Pt(
        0
    )

    instruction_heading.paragraph_format.space_after = Pt(
        6
    )

    title = instruction_title(
        target_name=target_name,
        target_kind=target_kind,
        object_name="obj_name",
        object_address="address_object",
    )

    _replace_paragraph_text(
        subtitle,
        title,
        font_size=TITLE_FONT_SIZE,
    )

    subtitle.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    subtitle.paragraph_format.space_before = Pt(
        0
    )

    subtitle.paragraph_format.space_after = Pt(
        12
    )


def _normalize_template_fonts(
    document: Document,
) -> None:
    """
    Принудительно приводит весь текст шаблона ИОТ
    к Times New Roman.
    """
    for paragraph in document.paragraphs:
        _set_paragraph_font(
            paragraph
        )

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _set_paragraph_font(
                        paragraph
                    )


def _format_body_paragraph(
    paragraph,
    text: str,
) -> None:
    is_heading = bool(
        SECTION_HEADING_RE.match(
            text
        )
        and not POINT_RE.match(
            text
        )
    )

    if is_heading:
        _replace_paragraph_text(
            paragraph,
            text,
            font_size=BODY_FONT_SIZE,
        )

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph.paragraph_format.first_line_indent = None

        # Единый заметный отступ между разделами.
        paragraph.paragraph_format.space_before = Pt(
            12
        )

        paragraph.paragraph_format.space_after = Pt(
            6
        )

        return

    _replace_paragraph_text(
        paragraph,
        text,
        font_size=BODY_FONT_SIZE,
    )

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.JUSTIFY
    )

    paragraph.paragraph_format.first_line_indent = Cm(
        1
    )

    paragraph.paragraph_format.space_before = Pt(
        0
    )

    paragraph.paragraph_format.space_after = Pt(
        3
    )


def _append_body(
    document: Document,
    instruction_text: str,
) -> None:
    lines = split_instruction_text(
        instruction_text
    )

    if not lines:
        raise RuntimeError(
            "Текст ИОТ пуст."
        )

    for line in lines:
        paragraph = document.add_paragraph()

        _format_body_paragraph(
            paragraph,
            line,
        )


def _replace_number_placeholders(
    document: Document,
) -> None:
    """
    В отдельной ИОТ оставляем /NMBP для ручной автозамены.

    Левую часть номера формируем до минут:
        DDMMYYHHMM/NMBP

    Секунды в номер отдельной инструкции не включаются.
    """
    number_prefix = datetime.now().strftime(
        "%d%m%y%H%M"
    )

    pattern = re.compile(
        r"(?:\d{10,14})?/NMBP"
    )

    replacement = (
        f"{number_prefix}/NMBP"
    )

    def process_paragraph(
        paragraph,
    ):
        text = paragraph.text

        if "/NMBP" not in text:
            return

        updated = pattern.sub(
            replacement,
            text,
        )

        if (
            updated == text
            and "/NMBP"
            in text
        ):
            updated = text.replace(
                "/NMBP",
                replacement,
                1,
            )

        if updated != text:
            _replace_paragraph_text(
                paragraph,
                updated,
            )

    for paragraph in document.paragraphs:
        process_paragraph(
            paragraph
        )

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    process_paragraph(
                        paragraph
                    )


def build_instruction_document(
    target_name: str,
    target_kind: str,
    instruction_text: str,
    output_path: str | Path,
    template_path: str | Path = TEMPLATE_PATH,
) -> Path:
    """
    Создаёт одну готовую ИОТ из шаблона.

    Ответственность этого модуля:
    - заголовок;
    - Times New Roman;
    - единые отступы;
    - тело инструкции;
    - номер DDMMYYHHMM/NMBP;
    - obj_name/address_object остаются тегами.

    Никакой работы с библиотекой, Telegram или VBA здесь нет.
    """
    template_path = Path(
        template_path
    )

    output_path = Path(
        output_path
    )

    if not template_path.exists():
        raise FileNotFoundError(
            "Не найден шаблон ИОТ: "
            f"{template_path}"
        )

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    document = Document(
        template_path
    )

    _normalize_template_fonts(
        document
    )

    _set_document_title(
        document,
        target_name=target_name,
        target_kind=target_kind,
    )

    _remove_old_body(
        document
    )

    _append_body(
        document,
        instruction_text,
    )

    _replace_number_placeholders(
        document
    )

    # Ещё один проход после всех замен,
    # чтобы вновь созданные runs тоже были Times New Roman.
    _normalize_template_fonts(
        document
    )

    document.save(
        output_path
    )

    return output_path


# Совместимость со старым кодом, если где-то ещё
# используется имя build_instruction().
def build_instruction(
    target_name: str,
    instruction_text: str,
    output_path: str | Path,
    target_kind: str = "position",
    template_path: str | Path = TEMPLATE_PATH,
) -> Path:
    return build_instruction_document(
        target_name=target_name,
        target_kind=target_kind,
        instruction_text=instruction_text,
        output_path=output_path,
        template_path=template_path,
    )
